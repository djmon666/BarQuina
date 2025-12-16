from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO

from flask import Blueprint, flash, redirect, render_template, url_for, send_file
from flask_login import current_user
from sqlalchemy import func
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from ...extensions import db
from ...models import CashSession, OrderItem, Payment

bp = Blueprint("reports", __name__, url_prefix="/reports")


@bp.before_request
def require_admin():
    from flask import current_app
    if current_app.config.get("LOGIN_DISABLED", False):
        return
    if not current_user.is_authenticated:
        flash("Cal iniciar sessió per accedir a aquesta pàgina.", "warning")
        return redirect(url_for("auth.login"))
    if not current_user.is_admin():
        flash("No tens permisos d'administrador.", "danger")
        return redirect(url_for("auth.login"))


@bp.route("/")
def index():
    """Display profitability and sales statistics per cash session."""
    sessions = CashSession.query.order_by(CashSession.opened_at.desc()).limit(20).all()

    session_stats = []
    for session in sessions:
        # Obtenir tots els OrderItems pagats durant aquesta sessió
        from ...models import Product, OrderItemExtra, PaymentItem
        
        paid_items = (
            db.session.query(OrderItem)
            .join(OrderItem.payment_links)
            .join(Payment)
            .filter(Payment.cash_session_id == session.id)
            .all()
        )
        
        # Agrupar productes amb els seus extres
        product_combinations = defaultdict(int)
        
        for item in paid_items:
            # Obtenir nom del producte
            product = db.session.get(Product, item.product_id)
            if not product:
                continue
            
            product_name = product.name
            
            # Afegir extres si n'hi ha
            if item.extras:
                extras_labels = sorted([extra.label for extra in item.extras])
                if extras_labels:
                    product_name += " (" + ", ".join(extras_labels) + ")"
            
            # Sumar quantitat
            product_combinations[product_name] += item.quantity
        
        # Convertir a llista ordenada per nom de producte
        products_sold = [
            {"name": name, "quantity": qty}
            for name, qty in sorted(product_combinations.items())
        ]

        session_stats.append(
            {
                "session": session,
                "products_sold": products_sold,
                "revenue": session.movement_net_total,
                "costs": session.inventory_costs,
                "profit": session.net_profit,
            }
        )

    return render_template("reports/index.html", session_stats=session_stats)


def _get_session_data(session_id):
    """Helper function to get session data for export."""
    from ...models import Product, OrderItemExtra, PaymentItem, InventoryEntry
    
    session = db.session.get(CashSession, session_id)
    if not session:
        return None
    
    # Obtenir productes venuts
    paid_items = (
        db.session.query(OrderItem)
        .join(OrderItem.payment_links)
        .join(Payment)
        .filter(Payment.cash_session_id == session.id)
        .all()
    )
    
    product_combinations = defaultdict(lambda: {"quantity": 0, "price": 0})
    for item in paid_items:
        product = db.session.get(Product, item.product_id)
        if not product:
            continue
        
        product_name = product.name
        total_price = product.price
        
        if item.extras:
            extras_labels = sorted([extra.label for extra in item.extras])
            if extras_labels:
                product_name += " (" + ", ".join(extras_labels) + ")"
            # Sumar preu dels extres
            for extra in item.extras:
                total_price += extra.price_delta * extra.quantity
        
        product_combinations[product_name]["quantity"] += item.quantity
        product_combinations[product_name]["price"] = total_price
    
    products_sold = [
        {"name": name, "quantity": data["quantity"], "price": data["price"]}
        for name, data in sorted(product_combinations.items())
    ]
    
    # Obtenir compres d'inventari
    inventory_entries = InventoryEntry.query.filter_by(cash_session_id=session.id).all()
    
    return {
        "session": session,
        "products_sold": products_sold,
        "revenue": session.movement_net_total,
        "costs": session.inventory_costs,
        "profit": session.net_profit,
        "inventory_entries": inventory_entries
    }


@bp.route("/export/pdf/<int:session_id>")
def export_pdf(session_id):
    """Export session report to PDF."""
    data = _get_session_data(session_id)
    if not data:
        flash("Sessió no trobada.", "danger")
        return redirect(url_for("reports.index"))
    
    session = data["session"]
    
    # Crear buffer per al PDF
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
    elements = []
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#2c3e50'),
        spaceAfter=20
    )
    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontSize=14,
        textColor=colors.HexColor('#34495e'),
        spaceAfter=10
    )
    
    # Títol
    opened = session.opened_at.strftime('%d/%m/%Y %H:%M') if session.opened_at else ''
    closed = session.closed_at.strftime('%d/%m/%Y %H:%M') if session.closed_at else 'Oberta'
    elements.append(Paragraph(f"Informe de Sessió #{session.id}", title_style))
    elements.append(Paragraph(f"Obertura: {opened} - Tancament: {closed}", styles['Normal']))
    elements.append(Spacer(1, 0.5*cm))
    
    # Productes venuts
    if data["products_sold"]:
        elements.append(Paragraph("Productes Venuts", heading_style))
        products_data = [["Producte", "Quantitat"]]
        for item in data["products_sold"]:
            products_data.append([item["name"], str(item["quantity"])])
        
        products_table = Table(products_data, colWidths=[12*cm, 3*cm])
        products_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3498db')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (1, 0), (1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey)
        ]))
        elements.append(products_table)
        elements.append(Spacer(1, 0.5*cm))
    
    # Resum financer
    elements.append(Paragraph("Resum Financer", heading_style))
    financial_data = [
        ["Concepte", "Import"],
        ["Ingressos (vendes)", f"{data['revenue']:.2f} €"],
        ["Costos (compres)", f"{data['costs']:.2f} €"],
        ["Benefici Net", f"{data['profit']:.2f} €"]
    ]
    
    financial_table = Table(financial_data, colWidths=[12*cm, 3*cm])
    financial_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2ecc71')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.lightgreen),
        ('BACKGROUND', (0, 3), (-1, 3), colors.HexColor('#27ae60')),
        ('TEXTCOLOR', (0, 3), (-1, 3), colors.whitesmoke),
        ('FONTNAME', (0, 3), (-1, 3), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey)
    ]))
    elements.append(financial_table)
    elements.append(Spacer(1, 0.5*cm))
    
    # Inventari
    if data["inventory_entries"]:
        elements.append(Paragraph("Compres d'Inventari", heading_style))
        inventory_data = [["Producte", "Quantitat", "Preu Unit", "Total"]]
        for entry in data["inventory_entries"]:
            inventory_data.append([
                entry.product_name,
                f"{entry.quantity:.2f}",
                f"{entry.unit_cost:.2f} €",
                f"{entry.total_cost:.2f} €"
            ])
        
        inventory_table = Table(inventory_data, colWidths=[8*cm, 2*cm, 2.5*cm, 2.5*cm])
        inventory_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e67e22')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#ffe4c4')),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey)
        ]))
        elements.append(inventory_table)
    
    # Generar PDF
    doc.build(elements)
    buffer.seek(0)
    
    filename = f"sessio_{session_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/pdf'
    )


@bp.route("/export/excel/<int:session_id>")
def export_excel(session_id):
    """Export session report to Excel."""
    data = _get_session_data(session_id)
    if not data:
        flash("Sessió no trobada.", "danger")
        return redirect(url_for("reports.index"))
    
    session = data["session"]
    
    # Crear workbook
    wb = Workbook()
    ws = wb.active
    ws.title = f"Sessió {session_id}"
    
    # Estils
    header_fill = PatternFill(start_color="3498db", end_color="3498db", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    title_font = Font(bold=True, size=14)
    
    # Informació de la sessió
    ws['A1'] = f"Informe de Sessió #{session.id}"
    ws['A1'].font = title_font
    
    opened = session.opened_at.strftime('%d/%m/%Y %H:%M') if session.opened_at else ''
    closed = session.closed_at.strftime('%d/%m/%Y %H:%M') if session.closed_at else 'Oberta'
    ws['A2'] = f"Obertura: {opened}"
    ws['A3'] = f"Tancament: {closed}"
    
    # Productes venuts
    row = 5
    ws[f'A{row}'] = "PRODUCTES VENUTS"
    ws[f'A{row}'].font = title_font
    row += 1
    
    ws[f'A{row}'] = "Producte"
    ws[f'B{row}'] = "Quantitat"
    ws[f'C{row}'] = "Preu (€)"
    ws[f'A{row}'].fill = header_fill
    ws[f'B{row}'].fill = header_fill
    ws[f'C{row}'].fill = header_fill
    ws[f'A{row}'].font = header_font
    ws[f'B{row}'].font = header_font
    ws[f'C{row}'].font = header_font
    row += 1
    
    for item in data["products_sold"]:
        ws[f'A{row}'] = item["name"]
        ws[f'B{row}'] = item["quantity"]
        ws[f'C{row}'] = item["price"]
        ws[f'C{row}'].number_format = '0.00'
        row += 1
    
    # Resum financer
    row += 2
    ws[f'A{row}'] = "RESUM FINANCER"
    ws[f'A{row}'].font = title_font
    row += 1
    
    ws[f'A{row}'] = "Concepte"
    ws[f'B{row}'] = "Import (€)"
    ws[f'A{row}'].fill = header_fill
    ws[f'B{row}'].fill = header_fill
    ws[f'A{row}'].font = header_font
    ws[f'B{row}'].font = header_font
    row += 1
    
    ws[f'A{row}'] = "Ingressos (vendes)"
    ws[f'B{row}'] = data['revenue']
    ws[f'B{row}'].number_format = '0.00'
    row += 1
    
    ws[f'A{row}'] = "Costos (compres)"
    ws[f'B{row}'] = data['costs']
    ws[f'B{row}'].number_format = '0.00'
    row += 1
    
    profit_fill = PatternFill(start_color="27ae60", end_color="27ae60", fill_type="solid")
    ws[f'A{row}'] = "Benefici Net"
    ws[f'B{row}'] = data['profit']
    ws[f'B{row}'].number_format = '0.00'
    ws[f'A{row}'].fill = profit_fill
    ws[f'B{row}'].fill = profit_fill
    ws[f'A{row}'].font = Font(bold=True, color="FFFFFF")
    ws[f'B{row}'].font = Font(bold=True, color="FFFFFF")
    row += 1
    
    # Inventari
    if data["inventory_entries"]:
        row += 2
        ws[f'A{row}'] = "COMPRES D'INVENTARI"
        ws[f'A{row}'].font = title_font
        row += 1
        
        ws[f'A{row}'] = "Producte"
        ws[f'B{row}'] = "Quantitat"
        ws[f'C{row}'] = "Preu Unit (€)"
        ws[f'D{row}'] = "Total (€)"
        for col in ['A', 'B', 'C', 'D']:
            ws[f'{col}{row}'].fill = header_fill
            ws[f'{col}{row}'].font = header_font
        row += 1
        
        for entry in data["inventory_entries"]:
            ws[f'A{row}'] = entry.product_name
            ws[f'B{row}'] = entry.quantity
            ws[f'B{row}'].number_format = '0.00'
            ws[f'C{row}'] = entry.unit_cost
            ws[f'C{row}'].number_format = '0.00'
            ws[f'D{row}'] = entry.total_cost
            ws[f'D{row}'].number_format = '0.00'
            row += 1
    
    # Ajustar amplada de columnes
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    
    # Guardar a buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    filename = f"sessio_{session_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
